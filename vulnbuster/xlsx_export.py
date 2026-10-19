# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Marcin Lewandowski <marcin.lewandowski.x@gmail.com>
"""Eksport raportów do Excela (openpyxl): podsumowanie + arkusz na każdą grupę (spółkę / zespół / typ…)."""
from __future__ import annotations

import io
import re
from datetime import date

from . import engine as E

COLS = [("Host", "hostname", 24), ("IP", "ip", 15), ("Spółka", "company", 18), ("Zespół", "team", 16),
        ("Typ", "type", 14), ("System", "os", 22), ("Środowisko", "environment", 11), ("Ekspozycja", "exposure", 11),
        ("Właściciel", "owner", 20), ("Administrator", "admin", 18), ("Aplikacja", "product", 22),
        ("Wersja", "version", 16), ("Docelowa", "target", 14), ("Status", "status", 14), ("Ryzyko", "risk", 11),
        ("Termin", "deadline", 11), ("Po terminie", "overdue", 10), ("KEV", "kev", 18), ("CVE", "cves", 30),
        ("Uzasadnienie", "reason", 60)]


def _sheet_name(label, used):
    n = re.sub(r"[\[\]:*?/\\]", " ", str(label or "—")).strip()[:28] or "—"
    base, i = n, 2
    while n.lower() in used:
        n = f"{base[:25]} ({i})"
        i += 1
    used.add(n.lower())
    return n


def grouped_xlsx(rep: dict, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")
    red = PatternFill("solid", fgColor="FDECEA")
    amber = PatternFill("solid", fgColor="FFF4DE")
    green = PatternFill("solid", fgColor="E3F5EA")
    wb = Workbook()
    ws = wb.active
    ws.title = "Podsumowanie"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Wygenerowano {date.today().isoformat()} · grupowanie: {rep['group_label']} · VulnBuster"
    ws["A2"].font = Font(italic=True, color="666666")
    hdr = [rep["group_label"], "Hosty", "Instalacje", "Do działań", "Podatności (otwarte)", "W KEV", "Po terminie",
           "Wyjątki", "Mitygowane", "Postęp %"]
    ws.append([])
    ws.append(hdr)
    for c in ws[4]:
        c.fill, c.font = head_fill, head_font
    for g in rep["groups"]:
        ws.append([g["label"], g["hosts"], g["installations"], g["open"], g["open_findings"], g["open_kev"],
                   g["overdue"], g["exception"], g["mitigated"], g["progress"]])
        row = ws.max_row
        ws.cell(row, 10).fill = green if g["progress"] >= 90 else amber if g["progress"] >= 60 else red
        if g["overdue"]:
            ws.cell(row, 7).font = Font(bold=True, color="B42318")
    t = rep["totals"]
    ws.append(["RAZEM", t["hosts"], t["installations"], t["open"], t["open_findings"], t["open_kev"], t["overdue"],
               t["exception"], t["mitigated"], t["progress"]])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for i, w in enumerate([30, 8, 11, 11, 18, 8, 12, 9, 11, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    used = {"podsumowanie"}
    for g in rep["groups"]:
        sh = wb.create_sheet(_sheet_name(g["label"], used))
        sh.append([c[0] for c in COLS])
        for c in sh[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="center")
        for r in g["rows"]:
            vals = []
            for _, key, _ in COLS:
                v = r.get(key)
                if key == "status":
                    v = E.LABELS.get(v, v)
                elif key == "risk":
                    v = E.RISK_LABELS.get(v, "") if v else ""
                elif key == "exposure":
                    v = E.EXPOSURE_LABELS.get(v, v) if v else ""
                elif key == "overdue":
                    v = "TAK" if v else ""
                elif key == "target":
                    v = v or ("migracja" if r.get("nofix") else "")
                elif isinstance(v, list):
                    v = ", ".join(v)
                vals.append(v if v is not None else "")
            sh.append(vals)
            row = sh.max_row
            st = r.get("status")
            fill = red if st == "vulnerable" else amber if st == "review" else None
            if fill:
                sh.cell(row, 14).fill = fill
            if r.get("overdue"):
                sh.cell(row, 16).font = Font(bold=True, color="B42318")
                sh.cell(row, 17).font = Font(bold=True, color="B42318")
        for i, (_, _, w) in enumerate(COLS, 1):
            sh.column_dimensions[get_column_letter(i)].width = w
        sh.freeze_panes = "B2"
        if sh.max_row > 1:
            sh.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{sh.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
