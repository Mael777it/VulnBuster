# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Marcin Lewandowski <marcin.lewandowski.x@gmail.com>
"""Import z Excela/CSV: parsowanie, automatyczne rozpoznawanie kolumn, zastosowanie do bazy."""
from __future__ import annotations

import base64
import csv
import fnmatch
import hashlib
import io
import re
import secrets
from collections import OrderedDict, defaultdict
from datetime import date, datetime

from . import engine as E
from . import versions as V
from .storage import now_iso

FIELDS = [
    ("ignore", "— pomiń —"),
    ("hostname", "Host / nazwa"),
    ("ip", "Adres IP"),
    ("company", "Spółka"),
    ("team", "Zespół / realizator"),
    ("host_type", "Typ urządzenia"),
    ("os", "System operacyjny"),
    ("os_version", "Wersja / build systemu"),
    ("environment", "Środowisko"),
    ("criticality", "Krytyczność"),
    ("exposure", "Ekspozycja (Internet / DMZ / wewn.)"),
    ("location", "Lokalizacja"),
    ("owner", "Właściciel / odpowiedzialny"),
    ("admin", "Administrator"),
    ("user", "Użytkownik"),
    ("product", "Aplikacja / produkt"),
    ("vendor", "Producent"),
    ("version", "Wersja zainstalowana"),
    ("cve", "CVE"),
    ("fixed_version", "Wersja naprawiona"),
    ("notes", "Uwagi"),
]
FIELD_IDS = {f for f, _ in FIELDS}
HOST_ATTRS = {"ip": "ip", "host_type": "type", "os": "os", "os_version": "os_version", "environment": "environment",
              "criticality": "criticality", "exposure": "exposure", "location": "location", "owner": "owner",
              "admin": "admin", "user": "user", "company": "company", "team": "team"}

SYN = {
    "hostname": ["host", "hostname", "nazwahosta", "serwer", "server", "maszyna", "komputer", "computer",
                 "computername", "device", "devicename", "urzadzenie", "nazwaurzadzenia", "asset", "assetname",
                 "zasob", "fqdn", "dnsname", "netbios", "nodename", "nazwaserwera", "servername", "hostnazwa"],
    "ip": ["ip", "adresip", "ipaddress", "address", "adres", "ipv4", "mgmtip", "ipadres"],
    "host_type": ["typ", "type", "rodzaj", "typhosta", "typurzadzenia", "devicetype", "assettype", "klasa",
                  "rodzajurzadzenia", "kategoriaurzadzenia", "assettypename", "typzasobu", "category", "kategoria"],
    "os": ["os", "system", "systemoperacyjny", "operatingsystem", "platform", "platforma", "osname", "ostype",
           "nazwasystemu"],
    "os_version": ["osversion", "osbuild", "build", "wersjasystemu", "osrelease", "kernel", "kernelversion",
                   "firmware", "firmwareversion", "wersjafirmware", "oprogramowaniesprzetowe", "versionos",
                   "buildsystemu", "osbuildnumber"],
    "company": ["spolka", "company", "firma", "podmiot", "organizacja", "organization", "organisation", "tenant",
                "companyname", "nazwaspolki", "spolkagrupy", "entity", "legalentity", "jednostka"],
    "team": ["zespol", "team", "realizator", "resolvergroup", "assignmentgroup", "supportgroup", "grupawsparcia",
             "zespolodpowiedzialny", "odpowiedzialnyzespol", "zespolrealizujacy", "patchteam"],
    "environment": ["srodowisko", "environment", "env", "stage"],
    "criticality": ["krytycznosc", "criticality", "waga", "istotnosc", "assetcriticality", "klasyfikacja"],
    "exposure": ["ekspozycja", "exposure", "wystawiony", "wystawionynazewnatrz", "internetfacing", "publiczny",
                 "dostepzinternetu", "strefa", "zone", "strefasieci", "networkzone", "zewnetrzny", "external"],
    "location": ["lokalizacja", "location", "site", "datacenter", "serwerownia", "oddzial", "lokacja"],
    "owner": ["wlasciciel", "owner", "odpowiedzialny", "responsible", "businessowner", "wlascicielbiznesowy",
              "opiekun", "osobaodpowiedzialna", "kontakt", "contact", "assetowner", "wlascicielaplikacji"],
    "admin": ["administrator", "admin", "technicalowner", "opiekuntechniczny", "adminsystemu"],
    "user": ["uzytkownik", "user", "username", "login", "lastuser", "lastloggeduser", "ostatniuzytkownik"],
    "product": ["aplikacja", "application", "app", "software", "produkt", "product", "program", "oprogramowanie",
                "komponent", "component", "pakiet", "package", "nazwaaplikacji", "softwarename", "usluga",
                "displayname", "nazwaoprogramowania"],
    "vendor": ["producent", "vendor", "dostawca", "manufacturer", "publisher", "wydawca", "softwarepublisher"],
    "version": ["wersja", "version", "wersjazainstalowana", "installedversion", "currentversion",
                "aktualnawersja", "ver", "wersjaaplikacji", "softwareversion", "zainstalowana"],
    "cve": ["cve", "cveid", "cves", "podatnosc", "podatnosci", "vulnerability", "vulnid", "identyfikatorcve"],
    "fixed_version": ["wersjanaprawiona", "fixedversion", "fixedin", "naprawionew", "naprawionawersja", "poprawkaw",
                      "patchedversion", "wersjapoprawiona", "wersjadocelowa", "targetversion", "fix",
                      "wersjazpoprawka", "minimalnawersja", "naprawiona", "poprawiona", "wersjabezpieczna"],
    "notes": ["uwagi", "notes", "komentarz", "comment", "opis", "description", "notatki", "notatka"],
}
_SYN_INDEX = [(syn, f) for f, lst in SYN.items() for syn in lst]

CVE_RE = re.compile(r"CVE-\d{4}-\d{4,}", re.I)
IP_RE = re.compile(r"^\d{1,3}(\.\d{1,3}){3}(/\d{1,2})?$")
VER_RE = re.compile(r"^(\d+:)?v?\d+([.\-_]\d+)+[\w.+~:()\-]*$|^v?\d+[a-z]{0,3}\d*$", re.I)
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", re.I)
HOST_RE = re.compile(r"^(?=.*[a-z])[a-z0-9][a-z0-9\-_]{1,62}(\.[a-z0-9\-_]+)*$", re.I)
COMBO_RE = re.compile(r"^(.+?)[\s_\-]+v?(\d+(?:\.\d+)+[\w.+~\-]*)$", re.I)
OS_KW = ["windows", "linux", "ubuntu", "debian", "rhel", "red hat", "centos", "rocky", "alma", "suse", "sles",
         "oracle linux", "ios-xe", "ios xe", "nx-os", "fortios", "junos", "pan-os", "esxi", "aix", "solaris",
         "macos", "freebsd", "cisco ios", "routeros", "win20", "server 20"]
TYPE_KW = ["serwer", "server", "switch", "przełącznik", "przelacznik", "router", "firewall", "workstation",
           "stacja", "laptop", "desktop", "vm", "wirtualna", "appliance", "ups", "drukarka", "printer",
           "load balancer", "storage", "hypervisor", "kontroler", "access point", "urządzenie sieciowe",
           "urzadzenie sieciowe", "sieciowe", "baza danych"]
ENV_KW = {"prod", "produkcja", "produkcyjne", "production", "test", "testowe", "dev", "uat", "preprod",
          "staging", "rozwojowe", "pre-prod", "szkoleniowe"}
CRIT_KW = {"wysoka", "średnia", "srednia", "niska", "high", "medium", "low", "critical", "krytyczna",
           "krytyczny", "wysoki", "niski", "średni"}


# ---------------------------------------------------------------- parsowanie wejścia
def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10]
    return str(v).strip()


def parse_table(text: str | None = None, xlsx_b64: str | None = None, sheet: str | None = None) -> dict:
    sheets, warnings = [], []
    if xlsx_b64:
        try:
            import openpyxl
        except ImportError:
            raise ValueError("Brak biblioteki openpyxl na serwerze — wklej dane przez schowek albo zapisz jako CSV")
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(xlsx_b64)), read_only=True, data_only=True)
        sheets = wb.sheetnames
        ws = wb[sheet] if sheet and sheet in sheets else wb[sheets[0]]
        rows, numeric_versions = [], 0
        for r in ws.iter_rows(values_only=True):
            row = [_cell_str(c) for c in r]
            numeric_versions += sum(1 for c in r if isinstance(c, float) and not c.is_integer())
            rows.append(row)
        if numeric_versions:
            warnings.append("Część komórek w Excelu to liczby (np. 14.1) — jeśli to wersje, mogły stracić końcowe "
                            "zera (14.10 → 14.1). Ustaw format kolumny na Tekst albo wklej przez schowek.")
    else:
        text = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip("\n")
        if not text.strip():
            raise ValueError("Brak danych do importu")
        first = text.split("\n", 5)[:5]
        if any("\t" in l for l in first):
            delim = "\t"
        else:
            delim = ";" if sum(l.count(";") for l in first) >= sum(l.count(",") for l in first) and any(
                ";" in l for l in first) else ","
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        raise ValueError("Brak danych do importu")
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    # usuń całkowicie puste kolumny
    keep = [i for i in range(width) if any(r[i] for r in rows)]
    rows = [[r[i] for i in keep] for r in rows]
    has_header = _looks_like_header(rows[0]) or len(rows) == 1
    if has_header:
        headers, body = rows[0], rows[1:]
        headers = [h or f"Kolumna {i + 1}" for i, h in enumerate(headers)]
    else:
        headers, body = [f"Kolumna {i + 1}" for i in range(len(rows[0]))], rows
    return {"headers": headers, "rows": body, "has_header": has_header, "sheets": sheets, "warnings": warnings}


def _looks_like_header(row) -> bool:
    if any(_header_score(c)[1] >= 0.75 for c in row if c):
        return True
    data_like = sum(1 for c in row if c and (CVE_RE.search(c) or IP_RE.match(c) or VER_RE.match(c)))
    return data_like == 0


# ---------------------------------------------------------------- rozpoznawanie kolumn
def _header_score(h):
    n = E.norm(h)
    if not n:
        return None, 0.0
    best, score = None, 0.0
    for syn, f in _SYN_INDEX:
        if n == syn:
            return f, 1.0
    for syn, f in _SYN_INDEX:
        if len(syn) >= 4 and syn in n:
            s = 0.6 + 0.25 * len(syn) / len(n)
            if s > score:
                best, score = f, s
    return best, score


def _value_scores(values):
    vals = [v for v in values if v][:300]
    if not vals:
        return {}
    n = len(vals)
    low = [v.lower() for v in vals]
    distinct = len(set(low)) / n
    f = lambda pred: sum(1 for v in vals if pred(v)) / n  # noqa: E731
    s = {}
    s["cve"] = 0.95 * f(lambda v: CVE_RE.search(v))
    s["ip"] = 0.9 * f(lambda v: IP_RE.match(v))
    s["version"] = 0.6 * f(lambda v: VER_RE.match(v) and not IP_RE.match(v))
    s["os"] = 0.75 * f(lambda v: any(k in v.lower() for k in OS_KW))
    s["host_type"] = 0.7 * f(lambda v: any(k == v.lower() or k in v.lower().split() for k in TYPE_KW))
    s["environment"] = 0.7 * f(lambda v: v.lower() in ENV_KW)
    s["criticality"] = 0.7 * f(lambda v: v.lower() in CRIT_KW)
    s["exposure"] = 0.6 * f(lambda v: any(k in v.lower() for k in ("internet", "dmz", "wewn", "zewn", "lan",
                                                                    "internal", "external", "public")))
    s["owner"] = 0.55 * f(lambda v: EMAIL_RE.match(v))
    hostish = f(lambda v: HOST_RE.match(v) and not VER_RE.match(v) and not any(k in v.lower() for k in OS_KW))
    digits = f(lambda v: any(ch.isdigit() for ch in v) or "-" in v or "." in v)
    s["hostname"] = hostish * (0.3 + 0.25 * distinct + 0.15 * digits)
    combo = f(lambda v: COMBO_RE.match(v) and not IP_RE.match(v))
    texty = f(lambda v: re.search(r"[a-ząćęłńóśźż]", v.lower()) and not CVE_RE.search(v) and not EMAIL_RE.match(v))
    s["product"] = max(0.5 * combo, texty * (0.4 - 0.2 * distinct))
    return s


def signature(headers) -> str:
    return hashlib.sha1("|".join(E.norm(h) for h in headers).encode()).hexdigest()[:16]


def analyze(db, table) -> dict:
    headers, rows = table["headers"], table["rows"]
    sig = signature(headers)
    profile = db.get("mapping_profiles", {}).get(sig) if table["has_header"] else None
    cols, cands = [], []
    for i, h in enumerate(headers):
        values = [r[i] for r in rows]
        hf, hs = _header_score(h) if table["has_header"] else (None, 0.0)
        vs = _value_scores(values)
        scores = defaultdict(float)
        for f, sc in vs.items():
            scores[f] = max(scores[f], sc)
        if hf:
            scores[hf] = max(scores[hf], hs)
            # nagłówek ma pierwszeństwo, ale wartości mogą go potwierdzić
            if vs.get(hf, 0) > 0.3:
                scores[hf] = min(1.0, scores[hf] + 0.1)
        for f, sc in scores.items():
            if sc >= 0.3:
                cands.append((sc, i, f))
        samples = [v for v in values if v][:4]
        cols.append({"index": i, "header": h, "samples": samples, "filled": sum(1 for v in values if v),
                     "field": "ignore", "confidence": 0.0, "source": None})
    cands.sort(key=lambda x: -x[0])
    used_cols, used_fields = set(), set()
    for sc, i, f in cands:
        if i in used_cols or (f in used_fields and f != "notes"):
            continue
        cols[i].update(field=f, confidence=round(sc, 2), source="nagłówek" if sc >= 0.6 else "zawartość")
        used_cols.add(i)
        used_fields.add(f)
    # druga kolumna z wersjami przy CVE -> wersja naprawiona
    if "cve" in used_fields and "fixed_version" not in used_fields:
        for c in cols:
            if c["field"] == "ignore" and _value_scores([r[c["index"]] for r in rows]).get("version", 0) > 0.4:
                c.update(field="fixed_version", confidence=0.4, source="zawartość")
                break
    if profile and len(profile["mapping"]) == len(cols):
        for c, f in zip(cols, profile["mapping"]):
            c.update(field=f, confidence=1.0, source="zapamiętane")
    combined = False
    pc = next((c for c in cols if c["field"] == "product"), None)
    if pc and not any(c["field"] == "version" for c in cols):
        vals = [r[pc["index"]] for r in rows if r[pc["index"]]][:100]
        combined = bool(vals) and sum(1 for v in vals if COMBO_RE.match(v)) / len(vals) > 0.5
    return {"columns": cols, "signature": sig, "profile_used": bool(profile),
            "fields": [{"id": f, "label": l} for f, l in FIELDS], "split_product_version": combined}


# ---------------------------------------------------------------- budowa rekordów
def build_records(table, mapping, split_combined=False):
    recs = []
    for ri, row in enumerate(table["rows"]):
        rec = {"_row": ri + (2 if table["has_header"] else 1)}
        for i, f in enumerate(mapping):
            if f == "ignore" or f not in FIELD_IDS or i >= len(row):
                continue
            v = row[i].strip()
            if not v:
                continue
            if f == "cve":
                rec.setdefault("cve", [])
                rec["cve"] += [c.upper() for c in CVE_RE.findall(v)]
            elif f == "notes" and "notes" in rec:
                rec["notes"] += " | " + v
            else:
                rec[f] = v
        if rec.get("exposure"):
            rec["exposure"] = normalize_exposure(rec["exposure"])
        if split_combined and rec.get("product") and not rec.get("version"):
            m = COMBO_RE.match(rec["product"])
            if m:
                rec["product"], rec["version"] = m.group(1).strip(), m.group(2)
        if len(rec) > 1:
            recs.append(rec)
    return recs


def normalize_exposure(v) -> str:
    s = E.norm(v)
    if not s:
        return ""
    if "dmz" in s:
        return "dmz"
    if any(k in s for k in ("internet", "zewn", "public", "extern", "publicz")):
        return "internet"
    if any(k in s for k in ("wewn", "intern", "lan", "prywat", "niewyst", "brak")) or s in ("nie", "no", "false", "0", "n"):
        return "internal"
    if s in ("tak", "yes", "true", "1", "t", "y"):
        return "internet"
    return "unknown"


# ---------------------------------------------------------------- normalizacja nazw aplikacji
_ARCH = re.compile(r"\s*[\(\[][^\)\]]*(x64|x86|amd64|arm64|64[- ]?bit|32[- ]?bit|\b[a-z]{2}-[a-z]{2}\b|"
                   r"x86_64|i386|i686|win64|win32)[^\)\]]*[\)\]]", re.I)
_TRAIL_DASH_VER = re.compile(r"\s+[-–]\s+v?\d+(\.\d+)+\S*\s*$")
_TRAIL_VER = re.compile(r"[\s_-]+v?\d+(\.\d+)*[a-z]?\s*$", re.I)
_TRAIL_WORDS = re.compile(r"[\s,]+(x64|x86|64-bit|32-bit|edition)\s*$", re.I)


def normalize_product_name(name: str) -> str:
    """'PostgreSQL 14' -> 'PostgreSQL', '7-Zip 23.01 (x64)' -> '7-Zip', 'Mozilla Firefox (x64 en-US)' -> 'Mozilla Firefox'."""
    n = re.sub(r"\s+", " ", str(name or "")).strip()
    if not n:
        return n
    prev = None
    while prev != n:
        prev = n
        n = _ARCH.sub("", n).strip()
        n = _TRAIL_DASH_VER.sub("", n).strip()
        n = _TRAIL_WORDS.sub("", n).strip()
        m = _TRAIL_VER.search(n)
        if m and len(n[: m.start()].strip()) >= 3:
            n = n[: m.start()].strip()
    return n.strip(" -–,") or str(name).strip()


_WIN_CLIENT = re.compile(r"^(microsoft\s+)?win(dows)?\s*(10|11|8\.1|7)\b", re.I)
_WIN_SERVER = re.compile(r"^(microsoft\s+)?win(dows)?\s*(server)?\s*(20(0[38]|1[2296]|2[2-5]))(\s*r2)?", re.I)
_OS_VER = re.compile(r"^(.*?)[\s_-]+v?(\d+(?:\.\d+)*)(?:\s+(lts|sp\d+))?\s*$", re.I)


def os_to_product(os_name: str, os_version: str = ""):
    """System operacyjny jako „aplikacja” do oceny CVE: (nazwa produktu, wersja)."""
    o = re.sub(r"\s+", " ", str(os_name or "")).strip()
    if not o:
        return None, None
    ver = str(os_version or "").strip()
    m = _WIN_SERVER.match(o)
    if m:
        return f"Windows Server {m.group(4)}{' R2' if m.group(6) else ''}", ver
    m = _WIN_CLIENT.match(o)
    if m:
        return f"Windows {m.group(3)}", ver
    m = _OS_VER.match(o)
    if m and len(m.group(1).strip()) >= 3:
        return m.group(1).strip(), ver or m.group(2)
    return o, ver


def compile_ignore(patterns: str):
    pats = [p.strip().lower() for p in str(patterns or "").split(",") if p.strip()]
    return lambda name: any(fnmatch.fnmatchcase(name.lower(), p) for p in pats)


# ---------------------------------------------------------------- zastosowanie
def new_id(prefix):
    return f"{prefix}_{secrets.token_hex(5)}"


def stage(db):
    """Płytka kopia do podglądu: wszystkie modyfikacje w apply_import są copy-on-write."""
    s = dict(db)
    for k in ("hosts", "products", "cves", "inventory", "mapping_profiles"):
        s[k] = dict(db[k])
    s["imports"] = list(db["imports"])
    return s


def _host_index(db):
    by_name, by_short, by_ip = {}, {}, {}
    for h in db["hosts"].values():
        n = (h.get("hostname") or "").lower()
        if n:
            by_name[n] = h["id"]
            by_short.setdefault(n.split(".")[0], h["id"])
        if h.get("ip"):
            by_ip.setdefault(h["ip"].strip(), h["id"])
    return by_name, by_short, by_ip


def product_index(db):
    idx = {}
    for p in db["products"].values():
        for n in [p["name"]] + list(p.get("aliases") or []):
            if "*" in n or "?" in n:
                continue
            if E.norm(n):
                idx.setdefault(E.norm(n), p["id"])
    return idx


def wildcard_aliases(db):
    """Aliasy z * / ? (np. „Mozilla Firefox*”) — dopasowanie po oryginalnej nazwie z importu."""
    out = []
    for p in db["products"].values():
        for a in p.get("aliases") or []:
            if "*" in a or "?" in a:
                out.append((a.strip().lower(), p["id"]))
    return out


def apply_import(db, records, name, opts=None):
    """Modyfikuje `db` (kopię z stage() albo prawdziwą bazę) i zwraca wynik z oceną wierszy."""
    opts = opts or {}
    ts = now_iso()
    imp_id = new_id("imp")
    by_name, by_short, by_ip = _host_index(db)
    pidx = product_index(db)
    new_hosts, new_products, rules_added, conflicts, warnings = [], [], [], [], []
    skipped = []
    flipped = set()
    obs = {}
    row_refs = []

    def get_host(rec):
        hn = (rec.get("hostname") or "").strip()
        ip = (rec.get("ip") or "").strip()
        hid = None
        if hn:
            hid = by_name.get(hn.lower()) or by_short.get(hn.lower().split(".")[0])
        if not hid and ip:
            hid = by_ip.get(ip)
        if hid:
            old = db["hosts"][hid]
            h = dict(old)
            if opts.get("update_hosts", True):
                for f, attr in HOST_ATTRS.items():
                    val = team_id(rec[f]) if f == "team" and rec.get(f) else rec.get(f)
                    if val and val != old.get(attr):
                        h[attr] = val
                if hn and not old.get("hostname"):
                    h["hostname"] = hn
            if h != old:
                h["updated_at"] = ts
                db["hosts"][hid] = h
            return hid
        if not hn and not ip:
            return None
        hid = new_id("h")
        h = {"id": hid, "hostname": hn or ip, "notes": "", "mitigations": [], "created_at": ts, "updated_at": ts}
        for f, attr in HOST_ATTRS.items():
            h[attr] = team_id(rec[f]) if f == "team" and rec.get(f) else rec.get(f, "")
        db["hosts"][hid] = h
        by_name[h["hostname"].lower()] = hid
        by_short.setdefault(h["hostname"].lower().split(".")[0], hid)
        if ip:
            by_ip.setdefault(ip, hid)
        new_hosts.append(h["hostname"])
        return hid

    settings = db.get("settings", {})
    ignored = compile_ignore(opts.get("ignore", settings.get("import_ignore", "")))
    normalize = opts.get("normalize", settings.get("import_normalize", True))
    unknown_mode = opts.get("unknown", "create")          # create | skip
    os_as_app = opts.get("os_as_app", settings.get("import_os_as_app", True))
    wild = wildcard_aliases(db)
    ignored_names, unknown_names, renamed = defaultdict(int), defaultdict(int), {}
    new_pids = {}
    team_ids = {E.norm(t["id"]): t["id"] for t in settings.get("teams", [])}
    team_ids.update({E.norm(t["name"]): t["id"] for t in settings.get("teams", [])})
    new_teams = []

    def team_id(v):
        n = E.norm(v)
        if not n:
            return ""
        if n not in team_ids:
            tid = n[:40]
            settings_teams = db["settings"].setdefault("teams", [])
            if not any(t["id"] == tid for t in settings_teams):
                db["settings"] = dict(db["settings"], teams=list(settings_teams) + [{"id": tid, "name": v.strip(),
                                                                                     "description": "z importu"}])
                new_teams.append(v.strip())
            team_ids[n] = tid
        return team_ids[n]

    def find_pid(raw, name):
        pid = pidx.get(E.norm(raw)) or pidx.get(E.norm(name))
        if not pid and wild:
            low = raw.lower()
            pid = next((p for a, p in wild if fnmatch.fnmatchcase(low, a)), None)
        return pid

    def get_product(rec):
        raw = (rec.get("product") or "").strip()
        if not raw:
            return None
        if ignored(raw):
            ignored_names[raw] += 1
            rec["_skip"] = "na liście ignorowanych"
            return None
        pn = normalize_product_name(raw) if normalize else raw
        if pn != raw:
            renamed[raw] = pn
        pid = find_pid(raw, pn)
        if not pid and unknown_mode == "skip":
            unknown_names[pn] += 1
            rec["_skip"] = "brak w rejestrze aplikacji"
            return None
        if pid:
            p = db["products"][pid]
            if p.get("status") != "used" and opts.get("mark_used", True):
                if p.get("status") == "not_used" and pid not in flipped:
                    flipped.add(pid)
                    warnings.append(f"„{p['name']}” był oznaczony jako NIEUŻYWANY, a import znalazł go na hostach — "
                                    f"status zmieniony na „używamy”. Sprawdź notatkę produktu.")
                p = dict(p, status="used", verified_at=ts[:10],
                         notes=((p.get("notes") or "") + f"\n[{ts[:10]}] Wykryty w imporcie — zmieniono z "
                                f"„{p.get('status')}” na „używamy”.").strip())
                db["products"][pid] = p
            return pid
        pid = new_id("p")
        db["products"][pid] = {"id": pid, "name": pn, "vendor": rec.get("vendor", ""), "aliases": [],
                               "category": rec.get("_category", ""), "status": "used", "notes": "",
                               "verified_at": ts[:10], "created_at": ts}
        pidx[E.norm(pn)] = pid
        new_products.append(pn)
        new_pids[pid] = pn
        return pid

    os_done = set()

    def add_rule(cid, pid, fixed):
        branch = V.derive_branch(fixed)
        old = db["cves"].get(cid)
        if old is None:
            cve = {"id": cid, "title": "", "description": "", "cvss": None, "groups": [], "rules": [],
                   "triage": {}, "notes": "", "created_at": ts, "updated_at": ts}
        else:
            cve = dict(old)
            cve["rules"] = [dict(r, branches=[dict(b) for b in r.get("branches", [])]) for r in old.get("rules", [])]
        rule = next((r for r in cve["rules"] if r["product_id"] == pid), None)
        if rule is None:
            rule = {"product_id": pid, "branches": [], "unlisted": "review"}
            cve["rules"].append(rule)
        ex = next((b for b in rule["branches"] if b.get("branch") == branch), None)
        if ex:
            if ex.get("fixed") and ex["fixed"] != fixed:
                conflicts.append(f"{cid} / gałąź {branch}: w bazie {ex['fixed']}, w imporcie {fixed} — bez zmian")
            return
        rule["branches"].append({"branch": branch, "introduced": "", "fixed": fixed, "state": "fix",
                                 "note": "z importu", "source": "import"})
        cve["updated_at"] = ts
        db["cves"][cid] = cve
        rules_added.append(f"{cid}: {db['products'][pid]['name']} {branch} → {fixed}")

    pairs = {}          # (hid, pid) -> [{"v": wersja, "c": [cve], "rec": wiersz}]

    def add_obs(hid, pid, ver, cves, rec):
        pairs.setdefault((hid, pid), []).append({"v": (ver or "").strip(), "c": list(cves or []), "rec": rec})

    for rec in records:
        hid = get_host(rec)
        if os_as_app and hid and rec.get("os") and hid not in os_done:
            os_done.add(hid)
            oname, over = os_to_product(rec["os"], rec.get("os_version", ""))
            if oname and not ignored(oname):
                orec = {"product": oname, "version": over, "_category": "system", "_row": rec["_row"]}
                saved = normalize
                normalize = False
                opid = get_product(orec)
                normalize = saved
                if opid:
                    add_obs(hid, opid, over, [], None)
        pid = get_product(rec)
        cves = rec.get("cve") or []
        if pid and cves and rec.get("fixed_version"):
            for cid in cves:
                add_rule(cid, pid, rec["fixed_version"])
        if hid and pid:
            add_obs(hid, pid, rec.get("version", ""), cves, rec)
        elif hid:
            row_refs.append((rec, None))
        elif pid and cves and rec.get("fixed_version"):
            row_refs.append((rec, "rule"))
        else:
            skipped.append(rec["_row"])

    # jedna instalacja na (host, aplikacja); kilka gałęzi naraz (np. PostgreSQL 14 i 16) = osobne instalacje
    def major(v):
        n = V.numeric_prefix(v)
        return str(n[0]) if n else ""

    def lowest(items):
        vs = [i["v"] for i in items if i["v"]]
        if not vs:
            return ""
        return sorted(vs, key=V.sort_key())[0]

    for (hid, pid), items in pairs.items():
        majors = OrderedDict()
        for it in items:
            majors.setdefault(major(it["v"]), []).append(it)
        if len([m for m in majors if m]) > 1:
            groups = [(f"{E.inv_key(hid, pid)}|{m or '?'}", grp) for m, grp in majors.items()]
        else:
            groups = [(E.inv_key(hid, pid), items)]
        for k, grp in groups:
            vers = {i["v"] for i in grp if i["v"]}
            v = lowest(grp)
            if len(vers) > 1:
                warnings.append(f"{db['hosts'][hid]['hostname']} / {db['products'][pid]['name']}: kilka wersji "
                                f"({', '.join(sorted(vers, key=V.sort_key()))}) — do oceny przyjęto najniższą {v}")
            cs = []
            for i in grp:
                cs += [c for c in i["c"] if c not in cs]
            obs[k] = {"k": k, "h": hid, "p": pid, "v": v, "c": cs}
            for i in grp:
                if i["rec"] is not None:
                    row_refs.append((i["rec"], k))

    # inwentarz
    hosts_in_import = set()
    existing = defaultdict(list)
    for k, inst in db["inventory"].items():
        existing[(inst["host_id"], inst["product_id"])].append(k)
    for k, o in obs.items():
        hosts_in_import.add(o["h"])
        old = db["inventory"].get(k)
        hist = list((old or {}).get("history", []))
        if not hist or hist[-1].get("version") != o["v"]:
            hist.append({"import_id": imp_id, "version": o["v"], "date": ts[:10]})
        db["inventory"][k] = {"key": k, "host_id": o["h"], "product_id": o["p"], "version": o["v"],
                              "reported_cves": o["c"], "first_seen": (old or {}).get("first_seen", ts),
                              "last_seen": ts, "import_id": imp_id, "removed": False, "history": hist[-6:]}
        # import jest wiążący dla pary host+aplikacja: inne „gałęzie” z poprzednich importów znikają
        for ok in existing.get((o["h"], o["p"]), []):
            if ok not in obs and not db["inventory"][ok].get("removed"):
                db["inventory"][ok] = dict(db["inventory"][ok], removed=True, removed_at=ts)
    removed = 0
    if opts.get("replace_scope"):
        for k, inst in list(db["inventory"].items()):
            if inst["host_id"] in hosts_in_import and k not in obs and not inst.get("removed"):
                db["inventory"][k] = dict(inst, removed=True, removed_at=ts)
                removed += 1

    db["imports"].append({"id": imp_id, "name": name or f"Import {ts[:16].replace('T', ' ')}",
                          "created_at": ts, "created_by": opts.get("user"), "rows": len(records), "observations": list(obs.values()),
                          "new_hosts": len(new_hosts), "new_products": len(new_products)})

    # ocena
    ctx = E.Ctx(db)
    inst_eval = {}
    for k, o in obs.items():
        fs = E.findings_for(ctx, db["hosts"][o["h"]], o["p"], o["v"], o["c"])
        inst_eval[k] = (fs, E.summarize(fs))
    rows_out = []
    for rec, k in row_refs:
        r = {"row": rec["_row"], "hostname": rec.get("hostname") or rec.get("ip") or "",
             "product": rec.get("product", ""), "version": rec.get("version", ""),
             "cve": ", ".join(rec.get("cve") or []), "status": "none", "reason": "", "target": None, "open": []}
        if k == "rule":
            r.update(status="rule", reason="Reguła wersji CVE (bez hosta)")
        elif k is None:
            if rec.get("_skip"):
                r.update(status="skip", reason=f"Aplikacja pominięta — {rec['_skip']} (dane hosta zaktualizowane)")
            else:
                r.update(status="host", reason="Tylko dane hosta")
        else:
            fs, s = inst_eval[k]
            mine = [f for f in fs if not rec.get("cve") or f["cve"] in rec["cve"]]
            r["status"] = E.worst(f["status"] for f in mine) if mine else "none"
            r["open"] = [f["cve"] for f in mine if f["status"] in E.ACTION]
            r["target"] = s["target"]
            r["reason"] = "; ".join(f"{f['cve']}: {f['reason']}" for f in mine[:3]) or "Brak znanych CVE dla produktu"
        rows_out.append(r)
    counts = defaultdict(int)
    for k, (fs, s) in inst_eval.items():
        counts[s["status"]] += 1
    per_new = defaultdict(int)
    for o in obs.values():
        if o["p"] in new_pids:
            per_new[o["p"]] += 1
    top = lambda d, n=150: [{"name": k, "count": v} for k, v in sorted(d.items(), key=lambda x: -x[1])[:n]]  # noqa: E731
    return {
        "import_id": imp_id,
        "stats": {"records": len(records), "installations": len(obs), "hosts": len(hosts_in_import),
                  "new_hosts": len(new_hosts), "new_products": len(new_products), "rules_added": len(rules_added),
                  "skipped": len(skipped), "removed": removed, "by_status": dict(counts)},
        "new_hosts": new_hosts[:200], "new_products": new_products[:200], "rules_added": rules_added[:200],
        "new_products_counts": sorted(({"name": new_pids[p], "count": c} for p, c in per_new.items()),
                                      key=lambda x: -x["count"])[:300],
        "ignored": top(ignored_names), "ignored_total": sum(ignored_names.values()),
        "unknown_skipped": top(unknown_names, 300), "unknown_total": sum(unknown_names.values()),
        "renamed": [{"from": a, "to": b} for a, b in list(renamed.items())[:60]], "renamed_total": len(renamed),
        "new_teams": new_teams,
        "conflicts": conflicts[:100], "warnings": warnings[:100],
        "skipped_rows": skipped[:50], "rows": rows_out[:2000], "rows_total": len(rows_out),
    }
